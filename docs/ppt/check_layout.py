#!/usr/bin/env python3
"""检查 deck 的版式问题：文字溢出容器、内容越出页面下边界。

没有 LibreOffice 时无法真正渲染 pptx，这里用「按字号估算文字渲染高度」的办法近似判断，
能抓住绝大多数溢出/越界问题（改了 build_deck.py 之后跑一次）。

用法：
    .venv/bin/python docs/ppt/check_layout.py docs/ppt/task_runner_overview.pptx

macOS 上还可以用 Quick Look 生成缩略图肉眼复核：
    qlmanage -t -s 1300 -o /tmp/out docs/ppt/task_runner_overview.pptx
"""
from __future__ import annotations

import math
import sys

from pptx import Presentation

EMU_IN = 914400.0
FOOTER_BAND_TOP = 6.9      # 页脚所在高度（in）：本就在最底部，不参与越界判定
CJK_START = 0x2E7F         # 大于此码点按全角宽度（1em）估算
W_LATIN = 0.52             # 拉丁字符宽度（em）
W_MONO = 0.60              # 等宽字体字符宽度（em）
LINE_FACTOR = 1.2          # 行高 = 字号 * 1.2 * line_spacing


def inch(emu) -> float:
    return (emu or 0) / EMU_IN


def _char_width(ch: str, size_pt: float, mono: bool) -> float:
    if mono:
        return W_MONO * size_pt
    return (1.0 if ord(ch) > CJK_START else W_LATIN) * size_pt


def est_text_height(tf) -> float:
    """估算文本帧的渲染高度（英寸）。"""
    total = 0.0
    for p in tf.paragraphs:
        runs = p.runs
        if not runs:
            total += 6 / 72.0
            continue
        size = max((r.font.size.pt if r.font.size else 18) for r in runs)
        text = "".join(r.text for r in runs)
        mono = any((r.font.name or "").lower().startswith("consolas") for r in runs)
        avail_in = inch(tf._parent.width) - inch(tf.margin_left) - inch(tf.margin_right)
        width_pt = sum(_char_width(c, size, mono) for c in text)
        lines = max(1, math.ceil(width_pt / max(avail_in * 72, 1) - 1e-6))
        spacing = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        total += lines * size * LINE_FACTOR * spacing / 72.0
        total += ((p.space_after.pt if p.space_after else 0)
                  + (p.space_before.pt if p.space_before else 0)) / 72.0
    return total + inch(tf.margin_top) + inch(tf.margin_bottom)


def main(path: str) -> int:
    prs = Presentation(path)
    limit = inch(prs.slide_height) - 0.05
    issues = 0

    for i, slide in enumerate(prs.slides, start=1):
        for sh in slide.shapes:
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            if inch(sh.top) > FOOTER_BAND_TOP:      # 页脚
                continue

            tf = sh.text_frame
            content_h = est_text_height(tf)
            top = inch(sh.top)
            box_h = inch(sh.height)
            bottom = top + max(box_h, content_h)
            filled = str(sh.fill.type) != "BACKGROUND (5)"
            label = tf.text.strip().replace("\n", " / ")[:46]

            if bottom > limit:
                print(f"[越界] P{i:02d} 底边 {bottom:.2f}in > 页面下限 {limit:.2f}in  <{label}>")
                issues += 1
            if filled and content_h > box_h + 0.04:
                print(f"[溢出] P{i:02d} 内容 {content_h:.2f}in > 容器 {box_h:.2f}in  <{label}>")
                issues += 1

    print(f"\n共 {len(prs.slides._sldIdLst)} 页，问题 {issues} 处")
    return 1 if issues else 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        raise SystemExit(2)
    raise SystemExit(main(sys.argv[1]))
